# -*- coding: utf-8 -*-
"""
pins.py + listings.json → Pinterest içe aktarma dosyaları.
    python3 sheet.py <import-template.xlsx> <çıktı_dizin>
Şablon: kullanıcının "pinterest-import-kit" zip'indeki import-template.xlsx
(satır 1 başlık, satır 2 yönlendirme metni, veri satır 3'ten). F/H/I açılır
listeleri ve J sütunundaki STATUS formülü korunur.

Görsel adresi: pinler repoda `public/pins/fw2627/` altında, panelin kendi
alan adından servis edilir (Higgsfield'a yükleme yok).

Link kuralı: `listings.json` her katalog id'si için Etsy listing numarasını ve
CANLI durumunu taşır (`ops/drafts-push?verify=1` ile okunur). Yalnız `active`
listing'e doğrudan link verilir; taslak listing ziyaretçiye açılmaz, o pin
mağaza sayfasına düşer ve çıktıda sayılır. Taslaklar yayına alınınca
listings.json yeniden okunur ve bu betik tekrar koşulur.
"""
import csv, json, os, re, sys
from collections import Counter
import openpyxl
from pins import PINS

tpl, out = sys.argv[1:3]
here = os.path.dirname(os.path.abspath(__file__))
listings = json.load(open(os.path.join(here, "listings.json")))
BASE = "https://amuletta.artifactstudio.info/pins/fw2627"
SHOP = "https://www.etsy.com/shop/byArtifactStudio"
BOARDS = {"Bracelets", "Necklaces", "Rings", "Studo Jewelry"}
N = len(PINS)

fallback = []


def link(key, n):
    if key == "SHOP":
        return SHOP
    row = listings[key]
    if row["state"] == "active":
        return f"https://www.etsy.com/listing/{row['listing_id']}"
    fallback.append((n, key))
    return SHOP


rows = []
for p in PINS:
    r = dict(image_url=f"{BASE}/pin-{p['n']:02d}.jpg", title=p["title"], description=p["desc"],
             alt_text=p["alt"], link=link(p["link"], p["n"]), board=p["board"], board_section="",
             priority=p["pr"], aspect_ratio="2:3")
    assert len(r["title"]) <= 100, (p["n"], len(r["title"]))
    assert len(r["description"]) <= 500 and len(r["alt_text"]) <= 500, p["n"]
    assert r["board"] in BOARDS and r["priority"] in ("low", "high")
    for k in ("title", "description", "alt_text"):
        assert not re.search("[–—â]", r[k]), (p["n"], k)  # em/en dash ve â yok
    rows.append(r)
assert len({r["image_url"] for r in rows}) == N
assert len({r["title"] for r in rows}) == N, "başlıklar tekil olmalı"

cols = ["image_url", "title", "description", "alt_text", "link", "board", "board_section", "priority", "aspect_ratio"]
wb = openpyxl.load_workbook(tpl)
ws = wb["Import Data"]
# Satır 2 şablonun yönlendirme metni; içe aktarma aracı onu pin sanıp
# reddeder (2026-09-26 dersi). Silinmez, boşaltılır; J formülleri kaymaz.
for c in range(1, ws.max_column + 1):
    ws.cell(row=2, column=c).value = None  # cell(value=None) atamayı ATLAR
for i, r in enumerate(rows):
    for j, c in enumerate(cols):
        ws.cell(row=3 + i, column=1 + j).value = r[c] or None
os.makedirs(out, exist_ok=True)
name = f"artifact-studio-fw2627-pinterest-{N}"
wb.save(os.path.join(out, name + ".xlsx"))
with open(os.path.join(out, name + ".csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.DictWriter(f, fieldnames=cols)
    w.writeheader()
    w.writerows(rows)
print(N, dict(Counter(r["board"] for r in rows)), dict(Counter(r["priority"] for r in rows)))
print("mağazaya düşen (taslak listing):", len(fallback), sorted({k for _, k in fallback}))
