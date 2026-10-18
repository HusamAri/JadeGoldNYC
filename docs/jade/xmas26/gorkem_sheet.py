"""Jade Gold NYC, Christmas 2026: the maker's cost sheet (Görkem) -> Jade-Noel-2026-Gorkem-Maliyet-Talebi.xlsx

Run: python3 -I docs/jade/xmas26/gorkem_sheet.py     (reads catalog.json + images/heroes/<ID>.jpg)

One row per piece: hero thumbnail, product type, name, the workshop dimensions and the enamel /
stone list from the catalog, our gold-gram estimate as a grey reference, and YELLOW cells for the
maker to fill: 14K gram, 14K $/g, labour, enamel, stone + setting, other, note. The total is a
formula over his cells only. No retail prices anywhere (maker book rule). Formulas stay within
what Numbers also knows (IF / N), so the file can be opened or saved as .numbers on a Mac.
Pattern: scripts/ophir/gen_cost_request_xlsx.py (2026-09-08).
"""
import json
import pathlib

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image

HERE = pathlib.Path(__file__).parent
OUT = HERE / "Jade-Noel-2026-Gorkem-Maliyet-Talebi.xlsx"
THUMBS = HERE / "build" / "thumbs"  # gitignored
THUMBS.mkdir(parents=True, exist_ok=True)

FONT = "Arial"
YELLOW = PatternFill("solid", fgColor="FFF2A8")
HDR = PatternFill("solid", fgColor="1F3864")
HDR_FILL = PatternFill("solid", fgColor="C00000")
REF = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
TYPE_TR = {"bracelet": "Bileklik", "necklace": "Kolye", "earrings": "Küpe", "brooch": "Broş"}
WEARER_TR = {"women": "kadın", "men": "erkek", "unisex": "unisex"}

cat = json.load(open(HERE / "catalog.json"))
products = cat["products"]
assert len(products) == 40

# thumbnails: 4:3 hero -> 240 x 180 px JPEG (40 x ~12 KB keeps the workbook small)
TW, TH = 240, 180
for p in products:
    src = HERE / "images" / "heroes" / f"{p['id']}.jpg"
    assert src.exists(), src
    im = Image.open(src).convert("RGB")
    im.thumbnail((TW, TH), Image.LANCZOS)
    im.save(THUMBS / f"{p['id']}.jpg", "JPEG", quality=82, optimize=True)

wb = Workbook()
ws = wb.active
ws.title = "Maliyet Talebi"
HEAD = [
    ("#", 4), ("Kod", 6), ("Görsel", 34), ("Ürün tipi", 10), ("Ürün adı", 26), ("Kime", 7),
    ("Ölçüler (atölye kartı)", 58), ("Mine", 22), ("Taş", 22), ("En büyük ölçü mm", 9),
    ("Bizim altın gram\ntahminimiz (14K)", 12),
    ("GRAM 14K\n(net, parça)", 11), ("14K $/g\n(her şey dahil ise)", 12), ("İŞÇİLİK USD\n(parça başı)", 12),
    ("MİNE USD", 10), ("TAŞ + MIHLAMA\nUSD", 13), ("DİĞER USD\n(döküm, cila…)", 12), ("ÜRETİCİ NOTU", 32),
    ("TOPLAM MALİYET\nUSD (otomatik)", 15),
]
FILL_COLS = range(12, 19)  # L..R are the maker's
ws.append([h for h, _ in HEAD])
for i, (h, w) in enumerate(HEAD, start=1):
    c = ws.cell(row=1, column=i)
    c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    c.fill = HDR_FILL if i in FILL_COLS else HDR
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 42

for n, p in enumerate(products, start=1):
    r = n + 1
    ws.row_dimensions[r].height = 140  # 180 px thumbnail ≈ 135 pt
    ws.cell(row=r, column=1, value=n)
    ws.cell(row=r, column=2, value=p["id"])
    img = XLImage(str(THUMBS / f"{p['id']}.jpg"))
    img.width, img.height = TW, TH
    ws.add_image(img, f"C{r}")
    ws.cell(row=r, column=4, value=TYPE_TR[p["category"]])
    ws.cell(row=r, column=5, value=p["name"])
    ws.cell(row=r, column=6, value=WEARER_TR.get(p["wearer"], p["wearer"]))
    ws.cell(row=r, column=7, value=p["workshop"]["dimensions"])
    ws.cell(row=r, column=8, value=p["enamel_colours_tr"] or "yok")
    ws.cell(row=r, column=9, value=p["stones_tr"] or "yok")
    ws.cell(row=r, column=10, value=p["max_mm"])
    ws.cell(row=r, column=11, value=p["gold_grams"]).fill = REF
    for col in FILL_COLS:
        ws.cell(row=r, column=col, value=None).fill = YELLOW
    ws.cell(row=r, column=19, value=f'=IF($L{r}="","",$L{r}*N($M{r})+N($N{r})+N($O{r})+N($P{r})+N($Q{r}))')
    for col in range(1, 20):
        cell = ws.cell(row=r, column=col)
        cell.border = BORDER
        cell.font = Font(name=FONT, size=9 if col == 7 else 10)
        cell.alignment = Alignment(wrap_text=col in (5, 7, 8, 9, 18), vertical="top" if col == 7 else "center",
                                   horizontal="center" if col in (1, 2, 4, 6, 10, 11, 12, 13) else "left")
        if col in (13, 14, 15, 16, 17, 19):
            cell.number_format = "$#,##0.00"
        if col in (11, 12):
            cell.number_format = "0.00"

ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:S{ws.max_row}"
note = ws.max_row + 2
ws.cell(row=note, column=3, value="SARI hücreler Görkem tarafından doldurulacak. Gram ve fiyat 14K SARI altın, parça başı net.").font = Font(name=FONT, bold=True, size=10, color="C00000")
ws.cell(row=note + 1, column=3, value="Her ölçü ve tezgâh sırası üretici kitabında (Jade-Noel-2026-Uretici-Kitabi.pdf); bu tablo yalnız maliyet içindir.").font = Font(name=FONT, italic=True, size=9)
ws.cell(row=note + 2, column=3, value="Bizim gram tahminimiz yalnız büyüklük fikri içindir; tartılmış değildir, referans alınmasın.").font = Font(name=FONT, italic=True, size=9)

# ---------------- Sheet 2: assumptions ----------------
a = wb.create_sheet("Varsayımlar")
a.column_dimensions["A"].width = 40
a.column_dimensions["B"].width = 14
a.column_dimensions["C"].width = 80
a["A1"] = "Varsayımlar / referanslar"
a["A1"].font = Font(name=FONT, bold=True, size=13)
rows = [
    ("Girdi", "Değer", "Kaynak / not"),
    ("14K $/g, önceki fatura (her şey dahil)", 106, "Önceki siparişlerde kullanılan oran: ham altın + işçilik + taş/boncuk dahil. Bu seri için GÜNCEL oranı SARI hücreye yazın."),
    ("10K $/g (sizin oranınız)", None, "SARI: panel 10K/18K varyantları da satıyor; gerçek oranınız yoksa boş bırakın, türetimle devam ederiz."),
    ("18K $/g (sizin oranınız)", None, "SARI: aynı."),
    ("Altın spot, referans (USD/ozt)", 4166.10, "gold-api.com, 2026-10-18 18:58 UTC. Sadece bilgi; fiyat sizin $/g'nizden kurulur."),
    ("Gram oranı 10K / 14K (aynı kalıp)", 0.892, "Alaşım yoğunluğu; 14K gramından panelde ölçeklenir."),
    ("Gram oranı 18K / 14K (aynı kalıp)", 1.137, "Alaşım yoğunluğu; 14K gramından panelde ölçeklenir."),
    ("Mine", "fırınlı tercih", "Çinko içermeyen 14K, ≤790 °C; olmazsa soğuk mine (epoksi), notta belirtin."),
    ("Taşlar", "lab-grown", "Küçük lab beyaz safir / renkli taşlar; taş + mıhlama ayrı sütunda."),
]
for i, (k, v, src) in enumerate(rows, start=3):
    a.cell(row=i, column=1, value=k).font = Font(name=FONT, bold=(i == 3), size=10)
    c = a.cell(row=i, column=2, value=v)
    c.font = Font(name=FONT, bold=True, size=10)
    if i in (4, 5, 6):
        c.fill = YELLOW
    if isinstance(v, (int, float)) and i in (4, 5, 6, 7):
        c.number_format = "$#,##0.00"
    a.cell(row=i, column=3, value=src).font = Font(name=FONT, italic=True, size=9)
    a.cell(row=i, column=3).alignment = Alignment(wrap_text=True, vertical="center")
a["A13"] = "TOPLAM MALİYET = gram × 14K $/g + işçilik + mine + taş/mıhlama + diğer. $/g her şey dahilse işçilik sütununu boş bırakın."
a["A13"].font = Font(name=FONT, italic=True, size=10)
a["A14"] = "Bu dosyada SATIŞ fiyatı YOK; önce gerçek maliyet toplanır, fiyat sonra panelde kurulur."
a["A14"].font = Font(name=FONT, bold=True, size=10, color="C00000")

# ---------------- Sheet 3: how to fill ----------------
h = wb.create_sheet("Nasıl Doldurulur")
h.column_dimensions["A"].width = 26
h.column_dimensions["B"].width = 100
h["A1"] = "Nasıl doldurulur"
h["A1"].font = Font(name=FONT, bold=True, size=13)
guide = [
    ("Amaç", "Noel 2026 serisinin 40 parçası (10 bileklik, 10 kolye, 10 küpe, 10 broş) için gerçek üretim maliyetini almak."),
    ("Doldurulacak", "Yalnız SARI hücreler: GRAM 14K, 14K $/g, İŞÇİLİK, MİNE, TAŞ + MIHLAMA, DİĞER, NOT. 'Varsayımlar' sayfasındaki üç sarı hücre (14K/10K/18K $/g) de sizin."),
    ("Gram", "14 ayar sarı altın, parça başı NET gram (zincir, kopça, küpe arkalığı, broş iğnesi dahil). Mine ve taş grama girmez."),
    ("$/g ya da işçilik", "Alışkanlığınız her şey dahil $/g ise onu yazın ve işçiliği boş bırakın. Parça başı işçilik veriyorsanız $/g'ye yalnız metali yazın."),
    ("Mine", "Fırınlı mine için parça başı bedel. Fırınlı olmuyorsa soğuk mine bedelini yazın ve NOT'a 'soğuk' deyin."),
    ("Taş + mıhlama", "Taş bedeli + mıhlama işçiliği toplamı. Taş yoksa 0."),
    ("Diğer", "Döküm, cila, rodaj, kaplama vb. varsa; yoksa 0."),
    ("Not", "Yapılamayan, değişmesi gereken ya da farklı ölçü öneren her şey buraya: ölçü, kalınlık, mine rengi, taş boyu."),
    ("Toplam", "Otomatik hesaplanır; elle yazmayın."),
    ("Görseller", "Her satırın görseli tasarım görselleştirmesidir (üretim fotoğrafı değil). Ölçü ve kesit için üretici kitabına bakın."),
]
r = 3
for k, v in guide:
    h.cell(row=r, column=1, value=k).font = Font(name=FONT, bold=True, size=10)
    c = h.cell(row=r, column=2, value=v)
    c.font = Font(name=FONT, size=10)
    c.alignment = Alignment(wrap_text=True, vertical="center")
    h.row_dimensions[r].height = 30
    r += 1

wb.save(OUT)
print("saved", OUT.name, f"{OUT.stat().st_size // 1024} KB", "rows:", len(products))
print("gram estimate checksum:", round(sum(i * p["gold_grams"] for i, p in enumerate(products, start=1)), 2))
